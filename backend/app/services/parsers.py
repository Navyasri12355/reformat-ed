"""Document parsers — extract page/slide text from PDF, DOCX, PPTX and TXT.

Each parser returns a list of ``{"page_num": int, "text": str}`` dicts. Parsing
libraries are imported lazily so the API boots even if an optional parser
dependency is missing; a clear error is raised only when that format is used.
"""

from __future__ import annotations

import io
import re
import unicodedata

SUPPORTED_TYPES = ("pdf", "docx", "pptx", "txt")

# Map common PDF-extraction artefacts to clean ASCII.
_REPLACEMENTS = {
    "‘": "'", "’": "'", "“": '"', "”": '"',
    "–": "-", "—": "-", "−": "-",
    "…": "...", " ": " ", "ﬁ": "fi", "ﬂ": "fl",
}
_BULLETS = re.compile(r"[•●◦▪‣·�]")


def detect_file_type(file_name: str) -> str | None:
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    return ext if ext in SUPPORTED_TYPES else None


def clean_text(text: str) -> str:
    """Normalise unicode, drop replacement/bullet glyphs, collapse whitespace."""
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", text)
    for bad, good in _REPLACEMENTS.items():
        text = text.replace(bad, good)
    text = _BULLETS.sub(" ", text)
    # Strip remaining non-printable control chars (keep newlines/tabs).
    text = "".join(ch for ch in text if ch == "\n" or ch == "\t" or not unicodedata.category(ch).startswith("C"))
    # Mend words hyphen-split across a line break:  "photo-\nsynthesis" -> "photosynthesis".
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    # Reflow PDF hard-wraps: a newline in the MIDDLE of a sentence (i.e. not
    # right after . ! ? and not a blank-line paragraph break) becomes a space,
    # so sentences are not cut at the column width.
    text = re.sub(r"(?<![.!?\n])[ \t]*\n(?!\n)[ \t]*", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse(file_bytes: bytes, file_type: str) -> list[dict]:
    if file_type == "pdf":
        pages = _parse_pdf(file_bytes)
    elif file_type == "docx":
        pages = _parse_docx(file_bytes)
    elif file_type == "pptx":
        pages = _parse_pptx(file_bytes)
    elif file_type == "txt":
        pages = _parse_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")
    for p in pages:
        p["text"] = clean_text(p.get("text", ""))
    return pages


def _parse_pdf(data: bytes) -> list[dict]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        pages.append({"page_num": i + 1, "text": text})
    return pages


def _parse_docx(data: bytes) -> list[dict]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # DOCX has no intrinsic pages — treat the whole document as one logical page.
    return [{"page_num": 1, "text": "\n".join(paragraphs)}]


def _parse_pptx(data: bytes) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    pages = []
    for i, slide in enumerate(prs.slides):
        chunks = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        chunks.append(line)
        pages.append({"page_num": i + 1, "text": "\n".join(chunks)})
    return pages


def _parse_txt(data: bytes) -> list[dict]:
    text = data.decode("utf-8", errors="replace").strip()
    return [{"page_num": 1, "text": text}]
